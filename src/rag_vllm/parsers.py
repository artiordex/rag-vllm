# =============================================================================
# 파일명: parsers.py
# 경로: src/rag_vllm/parsers.py
# 목적: 텍스트·PDF·DOCX·HWPX·이미지 문서 추출과 입력 한도 적용함
# 작성자: AI전략팀
# 작성일: 2026-09-30
# 수정일: 2026-09-30
# =============================================================================

"""텍스트·PDF·DOCX·HWPX·이미지 문서 추출과 입력 한도 적용함"""

from __future__ import annotations

import json
import mimetypes
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO
from pathlib import PurePath

from .chunking import MAX_TEXT_CHARACTERS, normalize_text

MAX_DOCUMENT_INPUT_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 2_048
MAX_ARCHIVE_MEMBER_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 2_000


class ParseError(ValueError):
    """파일을 안전하게 텍스트로 변환하지 못했음을 나타내는 예외임"""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    """추출 텍스트와 파서가 확인한 MIME·문서 메타데이터를 보관함"""

    text: str
    mime_type: str | None
    metadata: dict[str, object] = field(default_factory=dict)


class _HTMLTextExtractor(HTMLParser):
    """스크립트·스타일을 제외하고 HTML 본문 경계를 보존해 추출함"""

    def __init__(self) -> None:
        """HTML 텍스트 조각과 무시 중인 스크립트 깊이를 초기화함"""
        super().__init__()
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        """HTML 시작 태그에 따라 본문 경계 또는 무시 깊이를 갱신함"""
        if tag.lower() in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        elif tag.lower() in {"p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        """HTML 종료 태그에 따라 본문 경계를 추가하거나 무시 깊이를 줄임"""
        if tag.lower() in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        """스크립트·스타일 바깥의 텍스트 데이터만 수집함"""
        if not self._ignored_depth:
            self.parts.append(data)


def _check_archive_limits(archive: zipfile.ZipFile) -> None:
    """문서 ZIP 구성원 수와 압축 해제 크기로 폭탄형 입력을 차단함

    Raises:
        ParseError: 압축 파일이 구성원·단일 항목·전체 크기 한도를 초과할 때 발생함
    """
    members = archive.infolist()
    if len(members) > MAX_ARCHIVE_MEMBERS:
        raise ParseError("문서 압축 파일의 항목 수가 허용 한도를 초과했습니다.")
    if any(info.file_size > MAX_ARCHIVE_MEMBER_BYTES for info in members):
        raise ParseError("문서 압축 파일의 단일 항목 크기가 허용 한도를 초과했습니다.")
    if sum(info.file_size for info in members) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise ParseError("문서 압축 해제 후 크기가 허용 한도를 초과했습니다.")


def _with_parser_metadata(parsed: ParsedDocument, parser: str) -> ParsedDocument:
    """추출 결과를 정규화하고 파서·문자 수·상태 메타데이터를 덧붙임"""
    if len(parsed.text) > MAX_TEXT_CHARACTERS:
        raise ParseError(f"추출 텍스트가 허용 한도({MAX_TEXT_CHARACTERS}자)를 초과했습니다.")
    text = normalize_text(parsed.text)
    if len(text) > MAX_TEXT_CHARACTERS:
        raise ParseError(f"추출 텍스트가 허용 한도({MAX_TEXT_CHARACTERS}자)를 초과했습니다.")
    metadata = {
        **parsed.metadata,
        "parser": parser,
        "extraction_status": "success" if text else "empty",
        "extracted_characters": len(text),
    }
    return ParsedDocument(text=text, mime_type=parsed.mime_type, metadata=metadata)


def _decode_bytes(payload: bytes) -> str:
    """한국어 문서에서 자주 쓰이는 인코딩 순서로 바이트를 디코딩함"""
    for encoding in ("utf-8-sig", "cp949", "euc-kr", "utf-16", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _parse_pdf(payload: bytes) -> ParsedDocument:
    """PDF 텍스트와 표 행을 추출하고 페이지·문자 수 한도를 검증함

    Raises:
        ParseError: 의존성·페이지 수·추출 문자 수·PDF 구조가 허용 범위를
            벗어날 때 발생함
    """
    parser_name = "pypdf"
    try:
        import pdfplumber
    except ImportError:
        pdfplumber = None

    try:
        pages: list[str] = []
        extracted_characters = 0
        table_count = 0
        if pdfplumber is not None:
            parser_name = "pdfplumber"
            with pdfplumber.open(BytesIO(payload)) as pdf:
                page_count = len(pdf.pages)
                if page_count > MAX_PDF_PAGES:
                    raise ParseError(f"PDF 페이지 수가 허용 한도({MAX_PDF_PAGES})를 초과했습니다.")
                for page_number, page in enumerate(pdf.pages, start=1):
                    page_parts: list[str] = []
                    tables = page.find_tables() or []
                    table_regions = [table.bbox for table in tables]
                    text_page = page
                    if table_regions:
                        text_page = page.filter(
                            lambda item: item.get("object_type") != "char"
                            or not any(
                                x0 <= (item.get("x0", 0) + item.get("x1", 0)) / 2 <= x1
                                and top <= (item.get("top", 0) + item.get("bottom", 0)) / 2 <= bottom
                                for x0, top, x1, bottom in table_regions
                            )
                        )
                    page_text = normalize_text(text_page.extract_text(layout=True) or "")
                    if page_text:
                        page_parts.append(page_text)
                    for table in tables:
                        rows = [[str(cell or "").strip().replace("\n", " ") for cell in row] for row in table.extract()]
                        rows = [row for row in rows if any(row)]
                        if not rows:
                            continue
                        table_count += 1
                        width = max(len(row) for row in rows)
                        rows = [row + [""] * (width - len(row)) for row in rows]
                        page_parts.extend("| " + " | ".join(row) + " |" for row in rows[:1])
                        page_parts.append("| " + " | ".join("---" for _ in range(width)) + " |")
                        page_parts.extend("| " + " | ".join(row) + " |" for row in rows[1:])
                    combined_page = normalize_text("\n".join(page_parts))
                    if combined_page:
                        pages.append(f"[페이지 {page_number}]\n{combined_page}")
                        extracted_characters += len(combined_page)
                        if extracted_characters > MAX_TEXT_CHARACTERS:
                            raise ParseError("PDF에서 추출한 텍스트가 허용 한도를 초과했습니다.")
        else:
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(payload))
            page_count = len(reader.pages)
            if page_count > MAX_PDF_PAGES:
                raise ParseError(f"PDF 페이지 수가 허용 한도({MAX_PDF_PAGES})를 초과했습니다.")
            for page_number, page in enumerate(reader.pages, start=1):
                page_text = normalize_text(page.extract_text() or "")
                if page_text:
                    pages.append(f"[페이지 {page_number}]\n{page_text}")
                    extracted_characters += len(page_text)
                    if extracted_characters > MAX_TEXT_CHARACTERS:
                        raise ParseError("PDF에서 추출한 텍스트가 허용 한도를 초과했습니다.")

        is_scanned = page_count == 0 or (extracted_characters < 20 * page_count)
        pdf_metadata: dict[str, object] = {
            "page_count": page_count,
            "table_count": table_count,
            "is_scanned_pdf": is_scanned,
            "layout_parser": parser_name,
        }
        if is_scanned and not pages:
            from .multimodal import extract_visual_elements_from_pdf

            visual_elements = extract_visual_elements_from_pdf(payload)
            pdf_metadata["visual_elements_count"] = len(visual_elements)
            pages.append(
                f"[스캔 PDF 문서: 총 {page_count}페이지]\n"
                "- 텍스트 레이어가 없어 OCR 또는 멀티모달 분석이 필요합니다."
            )
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("PDF 문서를 읽을 수 없습니다.") from exc
    return ParsedDocument(
        text=normalize_text("\n\n".join(pages)),
        mime_type="application/pdf",
        metadata=pdf_metadata,
    )


def _parse_docx(payload: bytes) -> ParsedDocument:
    """DOCX 문단과 표 셀을 텍스트로 추출하고 압축 구조를 먼저 검증함

    Raises:
        ParseError: 의존성·압축 구조·문서 파싱·본문 크기 검증에 실패할 때 발생함
    """
    try:
        from docx import Document
    except ImportError as exc:  # NOTE: 선언된 DOCX 의존성이 테스트 환경에서 누락된 경우만 해당함
        raise ParseError("DOCX 파싱을 위해 python-docx를 설치해야 합니다.") from exc

    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            _check_archive_limits(archive)
        document = Document(BytesIO(payload))
    except zipfile.BadZipFile as exc:
        raise ParseError("손상되었거나 유효하지 않은 DOCX(ZIP) 파일입니다.") from exc
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("DOCX 문서를 읽을 수 없습니다.") from exc
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    table_count = 0
    for table in document.tables:
        table_count += 1
        table_rows: list[list[str]] = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                table_rows.append(cells)
        if table_rows:
            header = table_rows[0]
            parts.append("| " + " | ".join(header) + " |")
            parts.append("| " + " | ".join("---" for _ in header) + " |")
            for r in table_rows[1:]:
                parts.append("| " + " | ".join(r) + " |")
    return ParsedDocument(
        text=normalize_text("\n".join(parts)),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        metadata={"table_count": table_count},
    )


def _parse_hwpx(payload: bytes) -> ParsedDocument:
    """KS X 6101 HWPX에서 문단과 표를 추출함

    Raises:
        ParseError: HWPX 압축 구조나 XML 섹션이 유효하지 않을 때 발생함
    """
    import xml.etree.ElementTree as ET
    try:
        with zipfile.ZipFile(BytesIO(payload)) as zf:
            _check_archive_limits(zf)
            section_names = sorted(
                [name for name in zf.namelist() if "section" in name.lower() and name.endswith(".xml")]
            )
            if not section_names:
                raise ParseError("유효한 HWPX 섹션(section.xml)을 찾을 수 없습니다.")

            lines: list[str] = []
            table_count = 0

            for sec_name in section_names:
                xml_data = zf.read(sec_name)
                root = ET.fromstring(xml_data)

                def _lt(e: ET.Element) -> str:
                    """네임스페이스가 있는 HWPX 태그에서 로컬 이름만 추출함"""
                    return e.tag.split("}")[-1] if "}" in e.tag else e.tag

                # 표 셀(tc) 내부 요소들을 미리 집합으로 모아 일반 문단과 중복 추출 방지
                tc_descendants = set()
                for tc in root.iter():
                    if _lt(tc) == "tc":
                        for c in tc.iter():
                            if c is not tc:
                                tc_descendants.add(id(c))

                for elem in root.iter():
                    if id(elem) in tc_descendants:
                        continue
                    tag = _lt(elem)
                    if tag == "tbl":
                        table_count += 1
                        for tr in elem.iter():
                            if _lt(tr) == "tr":
                                cells: list[str] = []
                                for tc in tr:
                                    if _lt(tc) == "tc":
                                        cell_text = "".join(
                                            t.text or "" for t in tc.iter() if _lt(t) == "t" and t.text
                                        ).strip()
                                        cells.append(cell_text)
                                if any(cells):
                                    lines.append("| " + " | ".join(cells) + " |")
                    elif tag == "p":
                        if not any(_lt(c) == "tbl" for c in elem.iter()):
                            p_text = "".join(t.text or "" for t in elem.iter() if _lt(t) == "t" and t.text).strip()
                            if p_text:
                                lines.append(p_text)

            full_text = normalize_text("\n\n".join(lines))
            return ParsedDocument(
                text=full_text,
                mime_type="application/hwp+zip",
                metadata={"table_count": table_count, "section_count": len(section_names)},
            )
    except zipfile.BadZipFile as exc:
        raise ParseError("손상되었거나 유효하지 않은 HWPX(ZIP) 파일입니다.") from exc
    except Exception as exc:
        raise ParseError(f"HWPX 파싱 실패: {exc}") from exc


def _parse_xlsx(payload: bytes) -> ParsedDocument:
    """XLSX 스프레드시트의 시트별 표 데이터를 마크다운 표로 추출함

    Raises:
        ParseError: 압축 손상, 파일 크기 한도 초과 시 발생함
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ParseError("XLSX 파싱을 위해 openpyxl을 설치해야 합니다.") from exc

    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            _check_archive_limits(archive)
        wb = load_workbook(filename=BytesIO(payload), read_only=True, data_only=True)
    except zipfile.BadZipFile as exc:
        raise ParseError("손상되었거나 유효하지 않은 XLSX(ZIP) 파일입니다.") from exc
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"XLSX 문서를 읽을 수 없습니다: {exc}") from exc

    sheets_text: list[str] = []
    total_rows = 0
    sheet_names = wb.sheetnames

    try:
        for sheet_name in sheet_names:
            sheet = wb[sheet_name]
            rows: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() if c is not None else "" for c in row]
                if any(cells):
                    rows.append(cells)
                    total_rows += 1
                    if total_rows > 10_000:
                        break

            if not rows:
                continue

            sheet_lines = [f"[시트: {sheet_name}]"]
            max_cols = max(len(r) for r in rows)
            padded_rows = [r + [""] * (max_cols - len(r)) for r in rows]

            header = padded_rows[0]
            sheet_lines.append("| " + " | ".join(c.replace("\n", " ") for c in header) + " |")
            sheet_lines.append("| " + " | ".join("---" for _ in header) + " |")
            for row in padded_rows[1:]:
                sheet_lines.append("| " + " | ".join(c.replace("\n", " ") for c in row) + " |")

            sheets_text.append("\n".join(sheet_lines))
    finally:
        wb.close()
    return ParsedDocument(
        text=normalize_text("\n\n".join(sheets_text)),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        metadata={"sheet_count": len(sheet_names), "total_rows": total_rows},
    )


def _parse_csv(payload: bytes, delimiter: str = ",") -> ParsedDocument:
    """CSV 또는 TSV 데이터를 정규화된 마크다운 표로 변환함"""
    import csv

    raw_text = _decode_bytes(payload)
    reader = csv.reader(raw_text.splitlines(), delimiter=delimiter)
    rows: list[list[str]] = []
    for r in reader:
        if any(cell.strip() for cell in r):
            rows.append([cell.strip() for cell in r])
            if len(rows) > 10_000:
                break

    if not rows:
        return ParsedDocument(text="", mime_type="text/csv", metadata={"row_count": 0})

    max_cols = max(len(r) for r in rows)
    padded_rows = [r + [""] * (max_cols - len(r)) for r in rows]
    lines: list[str] = []
    lines.append("| " + " | ".join(c.replace("\n", " ") for c in padded_rows[0]) + " |")
    lines.append("| " + " | ".join("---" for _ in padded_rows[0]) + " |")
    for row in padded_rows[1:]:
        lines.append("| " + " | ".join(c.replace("\n", " ") for c in row) + " |")

    return ParsedDocument(
        text=normalize_text("\n".join(lines)),
        mime_type="text/csv" if delimiter == "," else "text/tab-separated-values",
        metadata={"row_count": len(rows), "col_count": max_cols},
    )


def _parse_hwp(payload: bytes) -> ParsedDocument:
    """레거시 HWP 5.0 OLE 복합 문서에서 BodyText 섹션 텍스트를 복구 및 추출함

    Raises:
        ParseError: OLE 구조가 손상되었거나 유효하지 않은 HWP일 때 발생함
    """
    import re
    import zlib

    try:
        import olefile
    except ImportError as exc:
        raise ParseError("HWP 파싱을 위해 olefile 라이브러리가 필요합니다.") from exc

    if not olefile.isOleFile(BytesIO(payload)):
        raise ParseError("유효한 OLE2 HWP 파일 형식이 아닙니다.")

    try:
        ole = olefile.OleFileIO(BytesIO(payload))
    except Exception as exc:
        raise ParseError(f"HWP OLE 파일을 열 수 없습니다: {exc}") from exc

    try:
        header_data = ole.openstream("FileHeader").read()
        is_compressed = bool(header_data[36] & 0x01) if len(header_data) >= 37 else True

        sections: list[str] = []
        for entry in ole.listdir():
            if len(entry) == 2 and entry[0] == "BodyText" and entry[1].startswith("Section"):
                stream_data = ole.openstream(entry).read()
                if is_compressed:
                    try:
                        decompressed = zlib.decompress(stream_data, -15)
                    except Exception:
                        try:
                            decompressed = zlib.decompress(stream_data)
                        except Exception:
                            decompressed = stream_data
                else:
                    decompressed = stream_data

                pos = 0
                sec_text_parts: list[str] = []
                while pos < len(decompressed):
                    if pos + 4 > len(decompressed):
                        break
                    header = int.from_bytes(decompressed[pos : pos + 4], "little")
                    pos += 4
                    tag_id = header & 0x3FF
                    size = (header >> 20) & 0xFFF
                    if size == 0xFFF:
                        if pos + 4 > len(decompressed):
                            break
                        size = int.from_bytes(decompressed[pos : pos + 4], "little")
                        pos += 4

                    if pos + size > len(decompressed):
                        break

                    record_data = decompressed[pos : pos + size]
                    pos += size

                    if tag_id == 67:  # HWPTAG_PARA_TEXT
                        try:
                            text_str = record_data.decode("utf-16le", errors="ignore")
                            cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text_str)
                            if cleaned.strip():
                                sec_text_parts.append(cleaned.strip())
                        except Exception:
                            continue

                if sec_text_parts:
                    sections.append("\n".join(sec_text_parts))

        full_text = normalize_text("\n\n".join(sections))
        if not full_text:
            raise ParseError("HWP 본문 텍스트를 추출하지 못했습니다. 손상되었거나 암호화된 문서일 수 있습니다.")

        return ParsedDocument(
            text=full_text,
            mime_type="application/x-hwp",
            metadata={"section_count": len(sections)},
        )
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"HWP 파싱 실패: {exc}") from exc
    finally:
        ole.close()


def parse_document(filename: str, payload: bytes, content_type: str | None = None) -> ParsedDocument:
    """지원 파일에서 텍스트를 추출하고 검색에 넣을 수 있는 형태로 반환함

    Args:
        filename: 업로드 파일명으로 확장자 판별에 사용함
        payload: 제한 검증을 거친 파일 바이트임
        content_type: 업로드가 제공한 MIME 타입임

    Returns:
        ParsedDocument: 정규화 텍스트, MIME 타입과 파서 메타데이터임

    Raises:
        ParseError: 파일 크기·형식·추출 결과가 지원 범위를 벗어날 때 발생함

    Caveats:
        스캔 PDF는 OCR 또는 멀티모달 분석 대상임을 메타데이터에 명시함
    """

    if len(payload) > MAX_DOCUMENT_INPUT_BYTES:
        raise ParseError(f"업로드 파일이 허용 한도({MAX_DOCUMENT_INPUT_BYTES}바이트)를 초과했습니다.")

    suffix = PurePath(filename or "").suffix.lower()
    mime_type = (content_type or mimetypes.guess_type(filename or "")[0] or "").split(";", 1)[0] or None

    if suffix == ".pdf" or mime_type == "application/pdf":
        return _with_parser_metadata(_parse_pdf(payload), "pdf")
    if suffix == ".docx":
        return _with_parser_metadata(_parse_docx(payload), "docx")
    if suffix == ".hwpx" or mime_type in {"application/hwp+zip", "application/haansofthwpx"}:
        return _with_parser_metadata(_parse_hwpx(payload), "hwpx")
    if suffix == ".hwp" or mime_type in {"application/x-hwp", "application/haansofthwp"}:
        return _with_parser_metadata(_parse_hwp(payload), "hwp")
    if suffix in {".xlsx", ".xlsm"} or mime_type in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    }:
        return _with_parser_metadata(_parse_xlsx(payload), "xlsx")
    if suffix == ".csv" or mime_type == "text/csv":
        return _with_parser_metadata(_parse_csv(payload, delimiter=","), "csv")
    if suffix == ".tsv" or mime_type == "text/tab-separated-values":
        return _with_parser_metadata(_parse_csv(payload, delimiter="\t"), "tsv")
    if suffix in {".html", ".htm"} or mime_type == "text/html":
        parser = _HTMLTextExtractor()
        parser.feed(_decode_bytes(payload))
        return _with_parser_metadata(ParsedDocument(normalize_text("".join(parser.parts)), "text/html"), "html")
    if suffix == ".json" or mime_type == "application/json":
        raw = _decode_bytes(payload)
        try:
            value = json.loads(raw)
            text = json.dumps(value, ensure_ascii=False, indent=2)
        except json.JSONDecodeError:
            text = raw
        return _with_parser_metadata(ParsedDocument(normalize_text(text), "application/json"), "json")
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}
    if suffix in IMAGE_EXTENSIONS or (mime_type and mime_type.startswith("image/")):
        from .multimodal import parse_image_document

        text, meta = parse_image_document(filename, payload, mime_type)
        return _with_parser_metadata(
            ParsedDocument(text=text, mime_type=mime_type or f"image/{suffix.lstrip('.')}", metadata=meta),
            "image",
        )

    return _with_parser_metadata(
        ParsedDocument(normalize_text(_decode_bytes(payload)), mime_type or "text/plain"),
        "text",
    )
